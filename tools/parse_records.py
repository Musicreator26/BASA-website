"""
Parse BASA national-record Excel files into the JSON format used by records.html.

Reads:
  NR 4Feb26/BRU NR (LC) updated 04February2026.xlsx
  NR 4Feb26/BRU NR (SC) updated 4February2026.xlsx

Writes:
  data/records-open.json
  data/records-age-group.json

Run from the project root:
  py tools/parse_records.py
"""
import openpyxl, json, re
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent

EVENT_MAP = {
    'free': 'Freestyle',
    'back': 'Backstroke',
    'breast': 'Breaststroke',
    'fly': 'Butterfly',
    'im': 'IM',
}

AGE_GROUP_MAP = {
    # Long course sheet names
    '7yrs': '7 years',
    '8yrs': '8 years',
    '9yrs': '9 years',
    '10yrs': '10 years',
    '11yrs': '11 years',
    '12-13yrs': '12-13 years',
    '14-15yrs': '14-15 years',
    '16-18yrs': '16-18 years',
    # Short course sheet names
    'Age Group 7 years': '7 years',
    'Age Group 8 years': '8 years',
    'Age Group 9 years': '9 years',
    'Age Group 10 years': '10 years',
    'Age Group 11 years': '11 years',
    'Age Group 12-13 years': '12-13 years',
    'Age Group 14-15 years': '14-15 years',
    'Age Group 16-18 years': '16-18 years',
}


def normalize_event(raw):
    s = str(raw).strip().lower()
    m = re.match(r'^(\d+)\s*(free|back|breast|fly|im)$', s)
    if m:
        return f"{m.group(1)}m {EVENT_MAP[m.group(2)]}"
    return str(raw).strip()


def normalize_time(t):
    if t is None:
        return None
    if isinstance(t, (int, float)):
        return f"{float(t):.2f}"
    s = str(t).strip()
    if ':' in s:
        return s
    parts = s.split('.')
    # "1.55.50" -> "1:55.50" (typo where minute/second separator was a dot)
    if len(parts) >= 3:
        minutes = parts[0]
        seconds = parts[1]
        hundredths = '.'.join(parts[2:])
        return f"{minutes}:{seconds}.{hundredths}"
    return s


def normalize_date(d):
    if d is None:
        return ""
    if isinstance(d, datetime):
        return d.strftime("%Y-%m-%d")
    return str(d).strip()


def parse_sheet(ws, course):
    """Yield record dicts from a single worksheet."""
    gender = None
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        while cells and cells[-1] is None:
            cells.pop()
        if not cells:
            continue

        first = cells[0]
        if first is None:
            continue
        first_s = str(first).strip().lower()

        # Gender section markers
        if first_s in ('boys', 'men'):
            gender = 'M'
            continue
        if first_s in ('girls', 'women'):
            gender = 'F'
            continue

        # Skip header rows
        if first_s == 'events':
            continue

        if len(cells) < 3:
            continue

        event = cells[0]
        swimmer = cells[1] if len(cells) > 1 else None
        time_val = cells[2] if len(cells) > 2 else None
        date_val = cells[3] if len(cells) > 3 else None
        venue = cells[4] if len(cells) > 4 else None

        # Skip vacant records
        if time_val is None or swimmer is None:
            continue
        if not gender:
            continue
        if not re.match(r'^\s*\d', str(event)):
            continue

        yield {
            'event': normalize_event(event),
            'gender': gender,
            'course': course,
            'time': normalize_time(time_val),
            'holder': str(swimmer).strip(),
            'location': str(venue).strip() if venue else '',
            'date': normalize_date(date_val),
        }


def main():
    files = [
        (ROOT / 'NR 4Feb26' / 'BRU NR (LC) updated 04February2026.xlsx', 'LCM'),
        (ROOT / 'NR 4Feb26' / 'BRU NR (SC) updated 4February2026.xlsx', 'SCM'),
    ]

    all_open = []
    all_age = []

    for path, course in files:
        wb = openpyxl.load_workbook(path, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sn_clean = sheet_name.strip()
            sn_lower = sn_clean.lower()

            # Skip relays and special-ability category for v1
            if 'relay' in sn_lower or 'special' in sn_lower:
                continue

            if sn_lower == 'open':
                all_open.extend(parse_sheet(ws, course))
            elif sn_clean in AGE_GROUP_MAP:
                ag = AGE_GROUP_MAP[sn_clean]
                for rec in parse_sheet(ws, course):
                    rec['age_group'] = ag
                    all_age.append(rec)

    open_obj = {
        "title": "Brunei National Open Records",
        "updated": "2026-02-04",
        "_instructions": "Add or edit records below. Required fields per record: event, gender (M or F), course (LCM=50m / SCM=25m), time (mm:ss.ss or ss.ss), holder, location, date. Save the file and refresh the website.",
        "records": all_open,
    }

    age_obj = {
        "title": "Brunei National Age Group Records",
        "updated": "2026-02-04",
        "_instructions": "Add or edit records below. Required fields: event, age_group, gender (M or F), course (LCM/SCM), time, holder, location, date.",
        "records": all_age,
    }

    (ROOT / 'data' / 'records-open.json').write_text(
        json.dumps(open_obj, indent=2, ensure_ascii=False), encoding='utf-8')
    (ROOT / 'data' / 'records-age-group.json').write_text(
        json.dumps(age_obj, indent=2, ensure_ascii=False), encoding='utf-8')

    # Summary
    print(f"Open records:      {len(all_open):4d}")
    print(f"Age-group records: {len(all_age):4d}")
    by_course_open = {}
    for r in all_open:
        by_course_open[r['course']] = by_course_open.get(r['course'], 0) + 1
    print(f"  Open by course: {by_course_open}")
    by_age = {}
    for r in all_age:
        by_age[r['age_group']] = by_age.get(r['age_group'], 0) + 1
    print(f"  Age groups: {dict(sorted(by_age.items()))}")


if __name__ == '__main__':
    main()
